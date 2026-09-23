"""
Generate one self-contained HTML slideshow per data class / topic.

Output files (written to the same directory as this script's PLOTS_ROOT parent):
  slideshow_bin_occupancy.html
  slideshow_bin_visit_threshold_timeseries.html
  slideshow_<class_label>.html   (one per data class found under plots/)
  slideshow_old.html             (if any archived plots exist)

Each file only embeds its own images, keeping file sizes manageable.

Keyboard navigation: left/right arrows.
Click sidebar item to jump.
Images fill the viewport; filename shown below.
"""

import re
from pathlib import Path
from collections import OrderedDict

PLOTS_ROOT  = Path('/Users/brucel/ecco/yip/profile_data/z_profile_file_analysis/plots')
OUTPUT_DIR  = Path('/Users/brucel/ecco/yip/profile_data/z_profile_file_analysis')

# Set True to embed all images as base64 — produces self-contained files
# that can be shared and opened anywhere.
EMBED_IMAGES = True

# Set True to also generate PowerPoint files.
BUILD_PPTX = False

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def label_for(path):
    stem = path.stem
    stem = re.sub(r'^(prob_map_ease2|prob_map|argo_bin|ncei)_?', '', stem)
    stem = stem.replace('_raw', '').replace('_', ' ')
    return stem.title()


THRESH_ORDER = {'1plus': 0, '2plus': 1, '3plus': 2}


def _prob_sort_key(path):
    bin_m    = re.search(r'_([ab])_', path.stem)
    thresh_m = re.search(r'(1plus|2plus|3plus)', path.stem)
    bin_ord  = 0 if (bin_m and bin_m.group(1) == 'a') else 1
    thresh_ord = THRESH_ORDER.get(thresh_m.group(1) if thresh_m else '', 9)
    return (bin_ord, thresh_ord)


def _detect_class_dirs():
    if not PLOTS_ROOT.exists():
        return []
    return sorted(
        d.name for d in PLOTS_ROOT.iterdir()
        if d.is_dir()
        and not d.name.startswith('old_')
        and (
            (d / 'bin_visit_probability_maps').exists()
            or (d / 'bin_visit_probability_maps_ease2').exists()
        )
    )

# ---------------------------------------------------------------------------
# Collect and categorise all PNGs
# Returns: OrderedDict[section_name, list of (Path, subsection_label|None)]
# ---------------------------------------------------------------------------

def collect_sections():
    class_dirs = _detect_class_dirs()
    class_set  = set(class_dirs)

    sections = OrderedDict()
    sections['Bin Occupancy'] = []
    sections['Visit Frequency Timeseries (>=N thresholds)'] = []
    for cls in class_dirs:
        sections[cls] = []
    sections['Old / Archived'] = []

    def add(section, path, subsection=None):
        sections[section].append((path, subsection))

    for p in sorted(PLOTS_ROOT.rglob('*.png')):
        rel   = p.relative_to(PLOTS_ROOT)
        parts = rel.parts

        if 'old_filetrees' in parts or parts[0].startswith('old_'):
            add('Old / Archived', p)
            continue

        top = parts[0]

        if top == 'bin_occupancy':
            add('Bin Occupancy', p)

        elif top == 'bin_visit_threshold_timeseries':
            add('Visit Frequency Timeseries (>=N thresholds)', p)

        elif top in class_set:
            rest = '/'.join(parts[1:])
            if 'bin_visit_probability_maps_ease2' in rest:
                map_type = 'EASE-Grid 2.0'
            elif 'bin_visit_probability_maps' in rest:
                map_type = 'Polygon'
            else:
                add('Old / Archived', p)
                continue

            is_per_year = 'per_year' in rest
            year_m = re.search(r'_(\d{4})_', p.stem)
            if is_per_year and year_m:
                subsec = f'{map_type} — {year_m.group(1)}'
            else:
                subsec = f'{map_type} — All Years'
            add(top, p, subsec)

        else:
            add('Old / Archived', p)

    def class_sort_key(item):
        path, subsec = item
        subsec = subsec or ''
        map_order = 0 if 'Polygon' in subsec else 1
        year_m    = re.search(r'(\d{4})', subsec)
        year      = int(year_m.group(1)) if year_m else 0
        bin_o, thr = _prob_sort_key(path)
        return (map_order, year, bin_o, thr)

    for cls in class_dirs:
        sections[cls].sort(key=class_sort_key)

    return OrderedDict((k, v) for k, v in sections.items() if v)

# ---------------------------------------------------------------------------
# HTML builder  (one section at a time)
# ---------------------------------------------------------------------------

def img_src(path):
    if EMBED_IMAGES:
        import base64
        with open(path, 'rb') as f:
            b64 = base64.b64encode(f.read()).decode('ascii')
        return f'data:image/png;base64,{b64}'
    else:
        return str(path.relative_to(OUTPUT_DIR))


def build_html(title, section_name, items):
    """
    items: list of (Path, subsection_label|None)
    Returns HTML string.
    """
    sidebar_items = []
    slides        = []
    total         = len(items)

    current_subsection = None

    # Single section — open it
    sidebar_items.append(
        f'<div class="sec-header" onclick="toggleSection(this)">'
        f'▾ {section_name}</div>'
        f'<div class="sec-items">'
    )

    for i, (path, subsec) in enumerate(items):
        if subsec and subsec != current_subsection:
            sidebar_items.append(f'<div class="subsec-header">{subsec}</div>')
            current_subsection = subsec

        lbl = label_for(path)
        sidebar_items.append(
            f'<div class="sidebar-item" onclick="goTo({i})" id="si-{i}">'
            f'{lbl}</div>'
        )

        if EMBED_IMAGES and (i + 1) % 20 == 0:
            print(f"    embedding {i+1}/{total}...")
        src = img_src(path)
        rel = str(path.relative_to(PLOTS_ROOT))
        slides.append({'label': lbl, 'section': section_name, 'rel': rel, 'src': src})

    sidebar_items.append('</div>')

    slides_js_str = '[' + ',\n'.join(
        f'{{"label": {repr(s["label"])}, "section": {repr(s["section"])}, '
        f'"src": {repr(s["src"])}, "rel": {repr(s["rel"])}}}'
        for s in slides
    ) + ']'

    sidebar_html = '\n'.join(sidebar_items)

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>{title}</title>
<style>
  * {{ box-sizing: border-box; margin: 0; padding: 0; }}
  body {{ display: flex; height: 100vh; font-family: sans-serif; background: #1a1a1a; color: #ddd; }}

  #sidebar {{
    width: 280px; min-width: 200px; max-width: 400px;
    background: #222; overflow-y: auto; flex-shrink: 0;
    border-right: 1px solid #444; resize: horizontal;
  }}
  #sidebar-title {{
    padding: 10px 12px; font-size: 13px; font-weight: bold;
    background: #333; border-bottom: 1px solid #555; color: #adf;
  }}
  .sec-header {{
    padding: 7px 10px; font-size: 12px; font-weight: bold;
    background: #2d2d2d; cursor: pointer; color: #9cf;
    border-top: 1px solid #444; user-select: none;
  }}
  .sec-header:hover {{ background: #3a3a3a; }}
  .sec-items {{ display: block; }}
  .sec-items.collapsed {{ display: none; }}
  .subsec-header {{
    padding: 4px 18px; font-size: 11px; color: #888;
    background: #252525; border-top: 1px solid #333;
    font-weight: bold; letter-spacing: 0.05em;
  }}
  .sidebar-item {{
    padding: 5px 24px; font-size: 11px; cursor: pointer;
    white-space: nowrap; overflow: hidden; text-overflow: ellipsis;
  }}
  .sidebar-item:hover {{ background: #383838; }}
  .sidebar-item.active {{ background: #1a4a7a; color: #fff; }}

  #viewer {{ flex: 1; display: flex; flex-direction: column; overflow: hidden; }}
  #toolbar {{
    display: flex; align-items: center; gap: 8px;
    padding: 6px 12px; background: #292929;
    border-bottom: 1px solid #444; flex-shrink: 0;
  }}
  #toolbar button {{
    padding: 4px 14px; background: #3a3a3a; color: #ddd;
    border: 1px solid #555; border-radius: 3px; cursor: pointer; font-size: 13px;
  }}
  #toolbar button:hover {{ background: #4a4a4a; }}
  #counter {{ font-size: 12px; color: #888; margin-left: 4px; }}
  #section-label {{ font-size: 12px; color: #9cf; margin-left: 8px; flex: 1; }}
  #img-label {{ font-size: 11px; color: #aaa; text-align: right; }}
  #img-container {{
    flex: 1; display: flex; align-items: center; justify-content: center;
    overflow: hidden; background: #111;
  }}
  #main-img {{ max-width: 100%; max-height: 100%; object-fit: contain; display: block; }}
  #filename {{
    padding: 4px 12px; font-size: 11px; color: #666;
    background: #1e1e1e; border-top: 1px solid #333;
    white-space: nowrap; overflow: hidden; text-overflow: ellipsis; flex-shrink: 0;
  }}
</style>
</head>
<body>
<div id="sidebar">
  <div id="sidebar-title">{title} ({len(slides)} plots)</div>
  {sidebar_html}
</div>
<div id="viewer">
  <div id="toolbar">
    <button onclick="prev()">&#8592; Prev</button>
    <button onclick="next()">Next &#8594;</button>
    <span id="counter"></span>
    <span id="section-label"></span>
    <span id="img-label"></span>
  </div>
  <div id="img-container"><img id="main-img" src="" alt=""></div>
  <div id="filename"></div>
</div>
<script>
const slides = {slides_js_str};
let idx = 0;
function goTo(i) {{
  if (i < 0) i = 0;
  if (i >= slides.length) i = slides.length - 1;
  idx = i;
  const s = slides[i];
  document.getElementById('main-img').src = s.src;
  document.getElementById('filename').textContent = s.rel;
  document.getElementById('counter').textContent = (i+1) + ' / ' + slides.length;
  document.getElementById('section-label').textContent = s.section;
  document.getElementById('img-label').textContent = s.label;
  document.querySelectorAll('.sidebar-item').forEach(el => el.classList.remove('active'));
  const si = document.getElementById('si-' + i);
  if (si) {{ si.classList.add('active'); si.scrollIntoView({{block: 'nearest'}}); }}
}}
function next() {{ goTo(idx + 1); }}
function prev() {{ goTo(idx - 1); }}
document.addEventListener('keydown', e => {{
  if (e.key === 'ArrowRight' || e.key === 'ArrowDown') next();
  if (e.key === 'ArrowLeft'  || e.key === 'ArrowUp')   prev();
}});
function toggleSection(header) {{
  const items = header.nextElementSibling;
  items.classList.toggle('collapsed');
  header.textContent = (items.classList.contains('collapsed') ? '▸ ' : '▾ ') +
                       header.textContent.slice(2);
}}
goTo(0);
</script>
</body>
</html>
"""

# ---------------------------------------------------------------------------
# PowerPoint (optional, one file per output)
# ---------------------------------------------------------------------------

def build_pptx(output_path, title, section_name, items):
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    SLIDE_W   = Inches(13.33)
    SLIDE_H   = Inches(7.5)
    BG_COLOR  = RGBColor(0x1a, 0x1a, 0x1a)
    TITLE_BG  = RGBColor(0x1a, 0x2a, 0x4a)
    SUBSEC_BG = RGBColor(0x22, 0x22, 0x2e)
    TEXT      = RGBColor(0xee, 0xee, 0xee)
    ACCENT    = RGBColor(0x88, 0xcc, 0xff)
    DIM       = RGBColor(0x88, 0x88, 0x88)

    def set_bg(slide, color):
        fill = slide.background.fill; fill.solid(); fill.fore_color.rgb = color

    def add_text(slide, text, left, top, w, h, size=18, bold=False,
                 color=TEXT, align=PP_ALIGN.LEFT):
        tf = slide.shapes.add_textbox(left, top, w, h).text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = align
        run = p.add_run(); run.text = text
        run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color

    prs = Presentation()
    prs.slide_width = SLIDE_W; prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank); set_bg(s, TITLE_BG)
    add_text(s, title, Inches(1), Inches(2.5), Inches(11.33), Inches(1.5),
             size=32, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
    add_text(s, f'{len(items)} plots', Inches(1), Inches(4.2), Inches(11.33),
             Inches(0.8), size=16, align=PP_ALIGN.CENTER)

    prev_subsec = None
    for path, subsec in items:
        if subsec and subsec != prev_subsec:
            s = prs.slides.add_slide(blank); set_bg(s, SUBSEC_BG)
            add_text(s, subsec, Inches(0.5), Inches(3.0), Inches(12.33), Inches(1.2),
                     size=24, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
            prev_subsec = subsec
        s = prs.slides.add_slide(blank); set_bg(s, BG_COLOR)
        s.shapes.add_picture(str(path), Emu(0), Inches(0.15), SLIDE_W, Inches(7.0))
        caption = label_for(path) + '   |   ' + str(path.relative_to(PLOTS_ROOT))
        add_text(s, caption, Inches(0.2), Inches(7.15), Inches(12.9),
                 Inches(0.3), size=9, color=DIM)

    prs.save(output_path)
    print(f"  saved: {output_path.name}")

# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

# Maps section name → output file stem
SECTION_TO_STEM = {
    'Bin Occupancy':                               'bin_occupancy',
    'Visit Frequency Timeseries (>=N thresholds)': 'bin_visit_threshold_timeseries',
    'Old / Archived':                              'old',
}


def main():
    print("Scanning plots...")
    sections = collect_sections()
    for k, v in sections.items():
        print(f"  {k}: {len(v)} images")

    for sec_name, items in sections.items():
        stem  = SECTION_TO_STEM.get(sec_name, sec_name)
        out_html = OUTPUT_DIR / f'slideshow_{stem}.html'
        title = f'Argo — {sec_name}'

        print(f"\nBuilding {out_html.name} ({len(items)} images)...")
        html = build_html(title, sec_name, items)
        out_html.write_text(html, encoding='utf-8')
        size_mb = out_html.stat().st_size / 1e6
        print(f"  saved: {out_html.name}  ({size_mb:.1f} MB)")

        if BUILD_PPTX:
            out_pptx = OUTPUT_DIR / f'slideshow_{stem}.pptx'
            build_pptx(out_pptx, title, sec_name, items)

    print("\nDone.")


if __name__ == '__main__':
    main()
